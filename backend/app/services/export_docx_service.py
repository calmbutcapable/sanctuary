from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from fastapi import HTTPException, status
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import or_, select
from sqlalchemy.orm import Session
from app.models.export_log import ExportLog

from app.models.thread import Thread
from app.models.plan import Plan
from app.models.step import Step
from app.models.user import User
from app.models.session_handoff import SessionHandoff

# Request / response-facing schemas

class ExportDocxRequest(BaseModel):
    include_progress_summary: bool = True
    include_completed_steps: bool = True
    include_project_notes: bool = True
    include_next_actions: bool = True

    include_reflection_summary: bool = False
    include_reflection_prompts: bool = False

    include_ai_support_notice: bool = True
    include_display_name: bool = False
    include_export_date: bool = True


class ExportLogCreate(BaseModel):
    thread_id: int
    user_id: str
    selected_options: dict
    status: str
    created_at: datetime
    error_message: str | None = None


# Internal export models used by the builder + generator

class ExportStepItem(BaseModel):
    position: int
    title: str
    description: str | None = None
    completed_at: datetime | None = None

# Minimal protocol-style repository layer.
# Replace these with your actual DB/service functions.

@dataclass
class ThreadRecord:
    id: int
    user_id: str
    title: str
    thread_type: str
    description: str | None
    created_at: datetime


@dataclass
class PlanRecord:
    id: int
    thread_id: int
    title: str
    description: str | None
    is_archived: bool
    created_at: datetime


@dataclass
class StepRecord:
    id: int
    plan_id: int
    title: str
    description: str | None
    position: int
    is_completed: bool
    completed_at: datetime | None
    created_at: datetime
    notes: str | None = None
    refined_output: str | None = None


@dataclass
class SessionHandoffRecord:
    id: int
    thread_id: int | None
    plan_id: int | None
    last_completed_text: str
    suggested_next_options: list[str]
    preferred_next_option: str | None
    is_active: bool
    created_at: datetime
    updated_at: datetime


@dataclass
class WeeklyReflectionRecord:
    total_entries: int
    active_days: int
    empty_days: int
    summary: str
    prompts: list[str]
    average_energy: float | None = None
    low_energy_checkins: int = 0


@dataclass
class ProfileRecord:
    id: str
    display_name: str


class ExportRepository:
    """
    Replace these stub methods with real database-backed queries.
    Keeping them together makes the route thin and the export logic testable.
    """

    def get_thread_for_user(self, thread_id: int, user_id: str) -> ThreadRecord | None:
        raise NotImplementedError
    
    def get_plan_for_user(self, plan_id: int, user_id: str) -> PlanRecord | None:
        raise NotImplementedError

    def get_active_plan_for_thread(self, thread_id: int) -> PlanRecord | None:
        raise NotImplementedError

    def get_steps_for_plan(self, plan_id: int) -> list[StepRecord]:
        raise NotImplementedError

    def get_active_handoff(self, thread_id: int, plan_id: int | None) -> SessionHandoffRecord | None:
        raise NotImplementedError

    def get_weekly_reflection_for_user(self, user_id: str) -> WeeklyReflectionRecord | None:
        raise NotImplementedError

    def get_profile(self, user_id: str) -> ProfileRecord | None:
        raise NotImplementedError

    def create_export_log(self, payload: ExportLogCreate) -> None:
        raise NotImplementedError

class ExportDocxData(BaseModel):
    model_config = ConfigDict(arbitrary_types_allowed=True)

    thread_title: str
    thread_description: str | None = None

    plan_title: str | None = None
    plan_description: str | None = None

    display_name: str | None = None
    export_date: str | None = None

    progress_summary: str | None = None
    most_recent_completed_step: str | None = None
    current_direction: str | None = None

    completed_steps: list[ExportStepItem] = Field(default_factory=list)
    project_notes: list[str] = Field(default_factory=list)
    next_actions: list[str] = Field(default_factory=list)
    reflection_summary: str | None = None
    reflection_activity_lines: list[str] = Field(default_factory=list)
    reflection_prompts: list[str] = Field(default_factory=list)

    ai_support_notice: str | None = None

class SqlAlchemyExportRepository(ExportRepository):
    def __init__(self, db: Session):
        self.db = db

    def get_thread_for_user(self, thread_id: int, user_id):
        thread = self.db.execute(
            select(Thread).where(
                Thread.id == thread_id,
                Thread.user_id == user_id,
            )
        ).scalars().first()

        if not thread:
            return None

        return ThreadRecord(
            id=thread.id,
            user_id=str(thread.user_id),
            title=thread.title,
            thread_type=thread.thread_type,
            description=thread.description,
            created_at=thread.created_at,
        )
    
    def get_plan_for_user(self, plan_id: int, user_id: str) -> PlanRecord | None:
        plan = self.db.execute(
            select(Plan)
            .join(Thread, Plan.thread_id == Thread.id)
            .where(
                Plan.id == plan_id,
                Thread.user_id == user_id,
            )
        ).scalars().first()

        if not plan:
            return None

        return PlanRecord(
            id=plan.id,
            thread_id=plan.thread_id,
            title=plan.title,
            description=plan.description,
            is_archived=plan.is_archived,
            created_at=plan.created_at,
        )

    def get_active_plan_for_thread(self, thread_id: int):
        plan = self.db.execute(
            select(Plan).where(
                Plan.thread_id == thread_id,
                Plan.is_archived == False
            ).order_by(Plan.created_at.desc())
        ).scalars().first()

        if not plan:
            return None

        return PlanRecord(
            id=plan.id,
            thread_id=plan.thread_id,
            title=plan.title,
            description=plan.description,
            is_archived=plan.is_archived,
            created_at=plan.created_at,
        )

    def get_steps_for_plan(self, plan_id: int):
        steps = self.db.execute(
            select(Step).where(
                Step.plan_id == plan_id
            ).order_by(Step.position.asc())
        ).scalars().all()

        return [
            StepRecord(
                id=s.id,
                plan_id=s.plan_id,
                title=s.title,
                description=s.description,
                position=s.position,
                is_completed=s.is_completed,
                completed_at=s.completed_at,
                created_at=s.created_at,
                notes=s.notes,
                refined_output=s.refined_output,
            )
            for s in steps
        ]

    def get_active_handoff(self, thread_id: int, plan_id: int | None):
        handoff = self.db.execute(
            select(SessionHandoff).where(
                SessionHandoff.is_active == True,
                or_(
                    SessionHandoff.thread_id == thread_id,
                    SessionHandoff.plan_id == plan_id,
                )
            ).order_by(SessionHandoff.updated_at.desc())
        ).scalars().first()

        if not handoff:
            return None

        return SessionHandoffRecord(
            id=handoff.id,
            thread_id=handoff.thread_id,
            plan_id=handoff.plan_id,
            last_completed_text=handoff.last_completed_text,
            suggested_next_options=handoff.suggested_next_options,
            preferred_next_option=handoff.preferred_next_option,
            is_active=handoff.is_active,
            created_at=handoff.created_at,
            updated_at=handoff.updated_at,
        )

    def get_weekly_reflection_for_user(self, user_id):
        return None

    def get_profile(self, user_id):
        user = self.db.execute(
            select(User).where(User.id == user_id)
        ).scalars().first()

        if not user:
            return None

        return ProfileRecord(
            id=str(user.id),
            display_name=user.email.split("@")[0],
        )

    def create_export_log(self, payload: ExportLogCreate) -> None:
        export_log = ExportLog(
        thread_id=payload.thread_id,
        user_id=payload.user_id,
        selected_options=payload.selected_options,
        status=payload.status,
        created_at=payload.created_at,
        error_message=payload.error_message,
    )
        self.db.add(export_log)
        self.db.commit()

# Export builder service

class DocxExportBuilder:
    AI_SUPPORT_NOTICE = (
        "This document was created with support from Sanctuary as an AI-assisted working draft. "
        "It is intended to support learning, planning, and reflection. "
        "This is not a final submission and should be reviewed, fact-checked, and personalised. "
        "The user remains responsible for all final content."
    )

    def _clean_text(self, value: str | None) -> str | None:
        if value is None:
            return None
        cleaned = value.strip()
        if not cleaned:
            return None
        if cleaned.lower() == "string":
            return None
        return cleaned
    
    def __init__(self, repository: ExportRepository):
        self.repository = repository

    def build_export_data(
        self,
        *,
        plan_id: int,
        current_user: User,
        options: ExportDocxRequest,
    ) -> ExportDocxData:
        plan = self.repository.get_plan_for_user(plan_id, current_user.id)
        if not plan:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Project not found",
            )

        thread = self.repository.get_thread_for_user(plan.thread_id, current_user.id)
        if not thread:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Project not found",
            )

        steps = self.repository.get_steps_for_plan(plan.id)
        handoff = self.repository.get_active_handoff(thread.id, plan.id)

        reflection = None
        if options.include_reflection_summary or options.include_reflection_prompts:
            reflection = self.repository.get_weekly_reflection_for_user(current_user.id)

        profile = None
        if options.include_display_name:
            profile = self.repository.get_profile(current_user.id)

        completed_steps = self._build_completed_steps(steps) if options.include_completed_steps else []
        progress_summary, most_recent_completed_step, current_direction = (None, None, None)
        if options.include_progress_summary:
            progress_summary, most_recent_completed_step, current_direction = self._build_progress_summary(
                steps=steps,
                handoff=handoff,
            )

        project_notes = self._build_project_notes(
            thread=thread,
            plan=plan,
            steps=steps,
            options=options,
        )
        next_actions = self._build_next_actions(steps=steps, handoff=handoff, options=options)
        reflection_summary, reflection_activity_lines = self._build_reflection_summary(reflection, options)
        reflection_prompts = self._build_reflection_prompts(reflection, options)

        return ExportDocxData(
            thread_title=self._clean_text(thread.title) or "Sanctuary Project",
            thread_description=self._clean_text(thread.description),
            plan_title=self._clean_text(plan.title) if plan else None,
            plan_description=self._clean_text(plan.description) if plan else None,
            display_name=profile.display_name if profile else None,
            export_date=datetime.now().strftime("%d %B %Y") if options.include_export_date else None,
            progress_summary=progress_summary,
            most_recent_completed_step=most_recent_completed_step,
            current_direction=current_direction,
            completed_steps=completed_steps,
            project_notes=project_notes,
            next_actions=next_actions,
            reflection_summary=reflection_summary,
            reflection_activity_lines=reflection_activity_lines,
            reflection_prompts=reflection_prompts,
            ai_support_notice=self.AI_SUPPORT_NOTICE if options.include_ai_support_notice else None,
        )

    def _build_completed_steps(self, steps: list[StepRecord]) -> list[ExportStepItem]:
        completed = [s for s in steps if s.is_completed]
        completed.sort(key=lambda item: item.position)
        return [
            ExportStepItem(
                position=step.position,
                title=step.title,
                description=step.description,
                completed_at=step.completed_at,
            )
            for step in completed
        ]

    def _build_progress_summary(
        self,
        *,
        steps: list[StepRecord],
        handoff: SessionHandoffRecord | None,
    ) -> tuple[str | None, str | None, str | None]:
        if not steps:
            return ("No project steps have been recorded yet.", None, None)

        total_steps = len(steps)
        completed_steps = [s for s in steps if s.is_completed]
        completed_count = len(completed_steps)

        latest_completed = None
        if completed_steps:
            latest_completed = max(
                completed_steps,
                key=lambda item: item.completed_at or item.created_at,
            )

        current_direction = None
        if handoff and handoff.preferred_next_option:
            current_direction = handoff.preferred_next_option
        else:
            next_incomplete = next((s for s in sorted(steps, key=lambda s: s.position) if not s.is_completed), None)
            if next_incomplete:
                current_direction = next_incomplete.title

        summary = f"{completed_count} of {total_steps} steps completed."
        return (
            summary,
            latest_completed.title if latest_completed else None,
            current_direction,
        )

    def _build_project_notes(
        self,
        *,
        thread: ThreadRecord,
        plan: PlanRecord | None,
        steps: list,
        options: ExportDocxRequest,
    ) -> list[str]:
        if not options.include_project_notes:
            return []

        notes: list[str] = []

        thread_desc = self._clean_text(thread.description)
        plan_desc = self._clean_text(plan.description) if plan else None

        if thread_desc:
            notes.append(thread_desc)

        if plan_desc and plan_desc not in notes and plan_desc != "Work through each step to gradually build your final project":
            notes.append(plan_desc)

   # 🔥 NEW: Add step notes
        for step in steps:
            if not hasattr(step, "notes"):
                continue

            section_lines: list[str] = []

            if step.refined_output:
                refined = self._clean_text(step.refined_output)
                if refined:
                    section_lines.append("Final Output:")
                    section_lines.append(refined)

            if step.notes:
                cleaned = self._clean_text(step.notes)
                if cleaned:
                    section_lines.insert(0, "Notes:")
                    section_lines.insert(1, cleaned)

            if section_lines:
                notes.append(step.title)

            if step.notes:
                cleaned = self._clean_text(step.notes)
                if cleaned:
                    notes.append(f"Notes: {cleaned}")

            if step.refined_output:
                refined = self._clean_text(step.refined_output)
                if refined:
                    notes.append(f"Final Output: {refined}")

            notes.append("")
            notes.extend(section_lines)
            notes.append("")  # spacing between steps

        return notes

    def _build_next_actions(
        self,
        *,
        steps: list[StepRecord],
        handoff: SessionHandoffRecord | None,
        options: ExportDocxRequest,
    ) -> list[str]:
        if not options.include_next_actions:
            return []

        actions: list[str] = []
        if handoff:
            if handoff.preferred_next_option:
                actions.append(handoff.preferred_next_option)
            for option in handoff.suggested_next_options:
                if option not in actions:
                    actions.append(option)

        if not actions:
            next_incomplete = next((s for s in sorted(steps, key=lambda s: s.position) if not s.is_completed), None)
            if next_incomplete:
                actions.append(next_incomplete.title)

        return actions

    def _build_reflection_summary(
        self,
        reflection: WeeklyReflectionRecord | None,
        options: ExportDocxRequest,
    ) -> tuple[str | None, list[str]]:
        if not options.include_reflection_summary:
            return (None, [])
        if not reflection:
            return ("No reflection summary is available for this export.", [])

        activity_lines = [
            f"Total entries: {reflection.total_entries}",
            f"Active days: {reflection.active_days}",
        ]
        return (reflection.summary, activity_lines)

    def _build_reflection_prompts(
        self,
        reflection: WeeklyReflectionRecord | None,
        options: ExportDocxRequest,
    ) -> list[str]:
        if not options.include_reflection_prompts:
            return []
        if not reflection or not reflection.prompts:
            return ["No reflection prompts are available for this export."]
        return reflection.prompts

# DOCX generator

class DocxGenerator:
    def generate(self, data: ExportDocxData) -> Path:
        document = Document()

        self._add_title_block(document, data)
        self._add_project_overview(document, data)
        self._add_progress_summary(document, data)
        self._add_completed_steps(document, data)
        self._add_project_notes(document, data)
        self._add_next_actions(document, data)
        self._add_reflection_summary(document, data)
        self._add_reflection_prompts(document, data)
        self._add_development_guidance(document)
        self._add_ai_support_notice(document, data)

        temp_file = NamedTemporaryFile(delete=False, suffix=".docx")
        path = Path(temp_file.name)
        temp_file.close()
        document.save(path)
        return path

    def _add_section_spacing(self, document: Document) -> None:
        document.add_paragraph("")

    def _add_title_block(self, document: Document, data: ExportDocxData) -> None:
        document.add_heading("Sanctuary Project Export", level=0)

        main_title = data.plan_title or data.thread_title

        p = document.add_paragraph()
        p.add_run("Project: ").bold = True
        p.add_run(main_title)

        if data.plan_title and data.thread_title and data.thread_title != data.plan_title:
            p = document.add_paragraph()
            p.add_run("Workspace: ").bold = True
            p.add_run(data.thread_title)

        if data.export_date:
            p = document.add_paragraph()
            p.add_run("Exported: ").bold = True
            p.add_run(data.export_date)

        if data.display_name:
            p = document.add_paragraph()
            p.add_run("Prepared for: ").bold = True
            p.add_run(data.display_name)

        self._add_section_spacing(document)

    def _add_project_overview(self, document: Document, data: ExportDocxData) -> None:
        overview_parts = []
        if data.thread_description:
            overview_parts.append(("Project description", data.thread_description))
        if data.plan_description:
            overview_parts.append(("Plan description", data.plan_description))

        if not overview_parts:
            return

        document.add_heading("Project Overview", level=1)
        for label, value in overview_parts:
            p = document.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(value)
            self._add_section_spacing(document)

    def _add_progress_summary(self, document: Document, data: ExportDocxData) -> None:
        if not data.progress_summary:
            return

        document.add_heading("Progress Summary", level=1)
        p = document.add_paragraph()
        p.add_run("Overall progress: ").bold = True
        p.add_run(data.progress_summary)

        if data.most_recent_completed_step:
            p = document.add_paragraph()
            p.add_run("Most recent completed step: ").bold = True
            p.add_run(data.most_recent_completed_step)

        if data.current_direction:
            p = document.add_paragraph()
            p.add_run("Current direction: ").bold = True
            p.add_run(data.current_direction)
            self._add_section_spacing(document)

    def _add_completed_steps(self, document: Document, data: ExportDocxData) -> None:
        if not data.completed_steps:
            return

        document.add_heading("Completed Steps", level=1)
        for step in data.completed_steps:
            p = document.add_paragraph()
            p.add_run(f"Step {step.position}: {step.title}").bold = True

            if step.description:
                p = document.add_paragraph()
                p.add_run("Description: ").bold = True
                p.add_run(step.description)

            if step.completed_at:
                p = document.add_paragraph()
                p.add_run("Completed: ").bold = True
                p.add_run(step.completed_at.strftime("%d %B %Y"))
            self._add_section_spacing(document)

    def _add_project_notes(self, document: Document, data: ExportDocxData) -> None:
        if not data.project_notes:
            return

        document.add_heading("Project Notes", level=1)

        current_title = None

        for note in data.project_notes:
            cleaned = note.strip()

            if not cleaned:
                continue

            if cleaned in {"Notes:", "Final Output:"}:
                p = document.add_paragraph()
                p.add_run(cleaned).bold = True
                continue

            if current_title is None or cleaned not in {"Notes:", "Final Output:"}:
                # Treat likely step titles as bold headings when they are short
                if len(cleaned) < 80 and not cleaned.endswith("."):
                    p = document.add_paragraph()
                    p.add_run(cleaned).bold = True
                    current_title = cleaned
                else:
                    document.add_paragraph(cleaned)

        self._add_section_spacing(document)

    def _add_next_actions(self, document: Document, data: ExportDocxData) -> None:
        if not data.next_actions:
            return

        document.add_heading("Next Actions", level=1)

        for action in data.next_actions:
            cleaned = str(action).strip()
            if cleaned:
                document.add_paragraph(cleaned, style="List Bullet")

        self._add_section_spacing(document)

    def _add_reflection_summary(self, document: Document, data: ExportDocxData) -> None:
        if not data.reflection_summary:
            return

        document.add_heading("Reflection Summary", level=1)
        p = document.add_paragraph()
        p.add_run("Summary: ").bold = True
        p.add_run(data.reflection_summary)

        if data.reflection_activity_lines:
            document.add_paragraph("Activity")
            for line in data.reflection_activity_lines:
                document.add_paragraph(line, style="List Bullet")
            self._add_section_spacing(document)

    def _add_reflection_prompts(self, document: Document, data: ExportDocxData) -> None:
        if not data.reflection_prompts:
            return

        document.add_heading("Reflection Prompts", level=1)
        for prompt in data.reflection_prompts:
            document.add_paragraph(prompt, style="List Bullet")
        self._add_section_spacing(document)

    def _add_ai_support_notice(self, document: Document, data: ExportDocxData) -> None:
        if not data.ai_support_notice:
            return

        document.add_heading("AI Support Notice", level=1)
        document.add_paragraph(data.ai_support_notice)
        self._add_section_spacing(document)

    def _add_development_guidance(self, document: Document) -> None:
        document.add_heading("Developing Your Project", level=1)

        document.add_paragraph(
            "This document is a starting point. You can improve and personalise it in ways that reflect your own thinking and understanding."
        )

        document.add_paragraph("You might consider:")

        points = [
            "Expanding each step with more detail about your decisions and why you made them",
            "Adding examples or real-world references to support your ideas",
            "Reflecting on what worked well and what you would change next time",
            "Adjusting the plan to better suit your own pace or approach",
            "Turning key steps into more detailed actions or outcomes",
        ]

        for point in points:
            document.add_paragraph(point, style="List Bullet")

        document.add_paragraph(
            "There is no single correct way to develop this project. The goal is to make it your own and show your understanding of the process."
        )

        self._add_section_spacing(document)
# Coordinator service

class ExportDocxService:
    def __init__(self, repository: ExportRepository, builder: DocxExportBuilder, generator: DocxGenerator):
        self.repository = repository
        self.builder = builder
        self.generator = generator

    def export_plan_docx(
        self,
        *,
        plan_id: int,
        current_user: User,
        options: ExportDocxRequest,
    ) -> Path:
        plan = self.repository.get_plan_for_user(plan_id, current_user.id)
        if not plan:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Project not found",
            )
        forced_options = options.model_copy(
                update={"include_ai_support_notice": True}
            )
        try:
            export_data = self.builder.build_export_data(
                plan_id=plan_id,
                current_user=current_user,
                options=forced_options,
            )
            output_path = self.generator.generate(export_data)
            self.repository.create_export_log(
                ExportLogCreate(
                    thread_id=plan.thread_id,
                    user_id=str(current_user.id),
                    selected_options=forced_options.model_dump(),
                    status="success",
                    created_at=datetime.utcnow(),
                )
            )
            return output_path
        except Exception as exc:  # noqa: BLE001
            print("DOCX EXPORT ERROR:", repr(exc))

            self.repository.create_export_log(
                ExportLogCreate(
                    thread_id=plan.thread_id,
                    user_id=str(current_user.id),
                    selected_options=forced_options.model_dump(),
                    status="failed",
                    created_at=datetime.utcnow(),
                    error_message=str(exc),
                )
            )

            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="The DOCX could not be generated right now. Please try again.",
            ) from exc